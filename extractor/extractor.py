import os
import json
import pdfplumber
import pytesseract
from PIL import Image
from docx import Document
import openpyxl
import requests
from kafka import KafkaProducer, KafkaConsumer

pytesseract.pytesseract.tesseract_cmd = r'C:\Program Files\Tesseract-OCR\tesseract.exe'  # Windows example

GOOGLE_API_KEY=""
GEMINI_MODEL="models/gemini-1.5-flash-latest"

GROQ_API_KEY=""
GROQ_MODEL="mixtral-8x7b-32768"

OPENROUTER_API_KEY=""
OPENROUTER_MODEL="mistralai/mistral-small-3.2-24b-instruct:free"

# Kafka Producer
producer = KafkaProducer(
    bootstrap_servers='localhost:9092',
    value_serializer=lambda v: json.dumps(v).encode('utf-8')
)

# Track processed document IDs in memory
processed_ids = set()

# -------- TEXT EXTRACTION -------- #
def extract_text(file_path):
    ext = file_path.lower().split(".")[-1]
    try:
        if ext == "pdf":
            with pdfplumber.open(file_path) as pdf:
                return "\n".join(page.extract_text() or "" for page in pdf.pages)
        elif ext in ["png", "jpg", "jpeg"]:
            return pytesseract.image_to_string(Image.open(file_path))
        elif ext == "docx":
            doc = Document(file_path)
            return "\n".join(p.text for p in doc.paragraphs)
        elif ext in ["xlsx", "xls"]:
            wb = openpyxl.load_workbook(file_path)
            text = ""
            for sheet in wb:
                for row in sheet.iter_rows(values_only=True):
                    text += "\t".join([str(cell) if cell else "" for cell in row]) + "\n"
            return text
        elif ext == "txt":
            with open(file_path, "r", encoding="utf-8") as f:
                return f.read()
        else:
            return ""
    except Exception as e:
        print(f"[Extractor ❌] Failed to extract from {file_path}: {e}")
        return ""

# -------- FORMATTERS -------- #
def fallback_format(text):
    return f"[Formatted]\n{text.strip()}"

def gemini_format(text):
    url = f"https://generativelanguage.googleapis.com/v1beta/models/gemini-2.0-flash:generateContent?key={GOOGLE_API_KEY}"
    payload = {
        "contents": [{"parts": [{"text": f"Clean this text:\n{text}"}]}]
    }
    try:
        r = requests.post(url, json=payload, timeout=10)
        r.raise_for_status()
        response_json = r.json()
        return response_json["candidates"][0]["content"]["parts"][0]["text"]
    except Exception as e:
        print(f"[Extractor ❌] Gemini formatting failed: {e}")
        return "[Unformatted] " + text

# -------- METADATA PROCESSING -------- #
def process_metadata(meta):
    doc_id = meta.get("document_id")
    if doc_id in processed_ids:
        print(f"[Extractor ⚠️] Skipping duplicate document: {doc_id}")
        return
    processed_ids.add(doc_id)

    path = meta.get("file_path") or meta.get("path")
    if not path or not os.path.exists(path):
        print(f"[Extractor ⚠️] File not found at path: {path}")
        return

    print(f"[Extractor 📥] Processing: {meta['document_name']} ({doc_id})")

    text = extract_text(path)
    if not text.strip():
        print(f"[Extractor ⚠️] No text extracted from: {path}")
        return

    try:
        formatted = fallback_format(text)
    except:
        try:
            formatted = gemini_format(text)
        except:
            formatted = "[Unformatted] " + text

    meta["extracted_text"] = formatted
    file_base = os.path.splitext(os.path.basename(path))[0]

    producer.send("doc.extracted", value=meta)
    print(f"[Extractor 📤] Sent to Kafka topic 'doc.extracted': {file_base}")

# -------- KAFKA CONSUMER -------- #
if __name__ == "__main__":
    consumer = KafkaConsumer(
        "doc.ingested",
        bootstrap_servers="localhost:9092",
        group_id="extractor-group",
        auto_offset_reset="latest",  # Only consume new messages
        enable_auto_commit=True,
        value_deserializer=lambda m: json.loads(m.decode("utf-8"))
    )

    print("[Extractor 🔄] Listening to Kafka topic 'doc.ingested'...")

    for message in consumer:
        metadata = message.value
        process_metadata(metadata)
